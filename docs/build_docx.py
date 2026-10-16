# -*- coding: utf-8 -*-
"""把《架构设计规范文档.md》渲染为 DOCX，并生成总体架构/部署拓扑图。"""
import os, re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"D:\DP_work_1\docs"
MD = os.path.join(BASE, "架构设计规范文档.md")
DOCX = os.path.join(BASE, "架构设计规范文档.docx")
ASSETS = os.path.join(BASE, "assets")
IMG = os.path.join(ASSETS, "架构与部署拓扑图.png")
os.makedirs(ASSETS, exist_ok=True)

FONT_CANDIDATES = [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyhbd.ttc", r"C:\Windows\Fonts\simhei.ttf"]

def load_font(size, bold=False):
    for p in ([FONT_CANDIDATES[1], FONT_CANDIDATES[0]] if bold else FONT_CANDIDATES):
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, size)
            except Exception:
                continue
    return ImageFont.load_default()

# ---------------- 1. 画总体架构与部署拓扑图 ----------------
def build_diagram():
    W, H = 1560, 1020
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    f_title = load_font(34, True)
    f_layer = load_font(22, True)
    f_box = load_font(20)
    f_small = load_font(17)

    d.text((W // 2, 30), "校园场馆预约与拼场组队系统  总体架构与部署拓扑", font=f_title, fill="#1F3864", anchor="ma")

    def rbox(x, y, w, h, lines, fill, outline, font=f_box, tc="#1a1a1a", radius=10):
        d.rounded_rectangle([x, y, x + w, y + h], radius=radius, fill=fill, outline=outline, width=2)
        if isinstance(lines, str):
            lines = [lines]
        total = len(lines) * (font.size + 6)
        cy = y + (h - total) // 2
        for ln in lines:
            d.text((x + w // 2, cy), ln, font=font, fill=tc, anchor="ma")
            cy += font.size + 6

    def layer_label(y, h, text, color):
        rbox(24, y, 130, h, text.split("\n"), color, color, font=f_layer, tc="white")

    def arrow(x, y1, y2, color="#7F7F7F"):
        d.line([x, y1, x, y2], fill=color, width=3)
        d.polygon([(x - 7, y2 - 10), (x + 7, y2 - 10), (x, y2)], fill=color)

    # 客户端层
    y, h = 100, 120
    layer_label(y, h, "客户端层", "#4472C4")
    names = ["学生端·微信小程序", "学生端·H5", "学生端·Android APK", "管理端·Web", "签到扫码终端"]
    bx, bw, gap = 178, 252, 20
    for i, n in enumerate(names):
        rbox(bx + i * (bw + gap), y, bw, h, n.split("·"), "#DEEAF6", "#4472C4", font=f_small)

    arrow(W // 2, y + h, y + h + 42)

    # 接入层
    y2, h2 = 262, 96
    layer_label(y2, h2, "接入层", "#BF8F00")
    rbox(178, y2, W - 178 - 30, h2, ["Nginx 1.24  —  HTTPS(443) / 静态资源托管 / /api 反向代理 / /ws 协议升级 / 跨域与限流"],
         "#FFF2CC", "#BF8F00")
    arrow(W // 2, y2 + h2, y2 + h2 + 42)

    # 应用层
    y3, h3 = 400, 230
    layer_label(y3, h3, "应用层", "#548235")
    rbox(178, y3, W - 178 - 30, 46, ["Spring Boot 3.2.12（模块化单体 · JDK 17 · MyBatis-Plus · Sa-Token）"],
         "#E2EFDA", "#548235", font=f_small)
    mods = ["认证鉴权\nSa-Token", "预约与\n槽位占用", "冲突检测\n防超卖", "拼场组队\n状态机", "候补递补\n队列",
            "签到与\n信用分", "消息中心\nWebSocket", "统计报表\nECharts"]
    mw, mg = 158, 12
    for i, m in enumerate(mods):
        rbox(186 + i * (mw + mg), y3 + 58, mw, 78, m.split("\n"), "#FFFFFF", "#548235", font=f_small, tc="#375623")
    rbox(186, y3 + 148, W - 186 - 38, 62,
         ["定时任务：未签到违约 · 候补超时顺延 · 消息补偿      |      AOP 操作日志 · Knife4j 接口文档 · Actuator 健康检查"],
         "#F2F7EE", "#A9D18E", font=f_small, tc="#375623")
    arrow(W // 2, y3 + h3, y3 + h3 + 42)

    # 数据层
    y4, h4 = 672, 118
    layer_label(y4, h4, "数据层", "#7030A0")
    dboxes = ["MySQL 8.0\n（唯一索引防超卖）", "Redis 7\n（可选·缓存与分布式锁）", "本地磁盘 / MinIO\n（可选·文件存储）"]
    dw = (W - 178 - 30 - 2 * 30) // 3
    for i, t in enumerate(dboxes):
        rbox(178 + i * (dw + 30), y4, dw, h4, t.split("\n"), "#EDE3F5", "#7030A0", font=f_small)

    # 右侧说明
    rbox(W - 20 - 0, 0, 0, 0, "", "white", "white") if False else None
    d.text((W // 2, 830), "部署要点：Ubuntu 22.04 (2C4G) + systemd 守护 jar + Nginx/HTTPS + 三网可访问；本地开发无需 Docker，使用原生 MySQL 8。",
           font=f_small, fill="#595959", anchor="ma")
    d.text((W // 2, 866), "降级策略：无 Redis → 唯一索引 + 本地缓存；无 MinIO → 本地磁盘；WebSocket 受限 → 5 秒轮询兜底。",
           font=f_small, fill="#595959", anchor="ma")
    d.text((W // 2, 902), "多端互动：业务变更先落库、再推送，msgId 去重，重连后按 lastId 补齐，保证消息不丢。",
           font=f_small, fill="#595959", anchor="ma")
    img.save(IMG)
    return IMG

# ---------------- 2. Markdown -> DOCX ----------------
def set_fonts(run, ascii_font="Times New Roman", ea_font="宋体", size=10.5, bold=None, color=None):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:eastAsia'), ea_font)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")

def add_inline(par, text, ascii_font="Times New Roman", ea_font="宋体", size=10.5, base_bold=False, color=None):
    for seg in INLINE.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = par.add_run(seg[2:-2]); set_fonts(r, ascii_font, ea_font, size, True, color)
        elif seg.startswith("`") and seg.endswith("`"):
            r = par.add_run(seg[1:-1]); set_fonts(r, "Consolas", "宋体", size - 0.5, base_bold, "C00000" if not color else color)
        else:
            r = par.add_run(seg); set_fonts(r, ascii_font, ea_font, size, base_bold, color)

def shade(el, fill):
    tcPr = el.get_or_add_tcPr() if el.tag.endswith('}tc') else el.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_table_fixed(table, widths_cm):
    """强制固定表格布局并写入 tblGrid，保证 Word/LibreOffice 都按设计宽度渲染。"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    table.autofit = False
    # 固定布局
    for tag in ('w:tblLayout', 'w:tblW'):
        el = tblPr.find(qn(tag))
        if el is not None:
            tblPr.remove(el)
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
    total_tw = sum(int(round(w * 567)) for w in widths_cm)   # 1cm = 566.93 twips
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(total_tw)); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc in list(grid):
            grid.remove(gc)
        for w in widths_cm:
            gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(round(w * 567)))); grid.append(gc)
    return total_tw

def build_docx(diagram):
    with open(MD, encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = "Times New Roman"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)

    i, n = 0, len(lines)
    img_inserted = False
    while i < n:
        line = lines[i].rstrip()
        # 代码块
        if line.startswith("```"):
            i += 1; buf = []
            while i < n and not lines[i].startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.left_indent = Cm(0.4)
            if buf:
                r = p.add_run("\n".join(buf)); set_fonts(r, "Consolas", "宋体", 8.5)
            shade(p._p, "F4F4F4")
            continue
        # 表格
        if line.startswith("|"):
            block = []
            while i < n and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            rows = []
            for b in block:
                if re.fullmatch(r"\|[\s:\-\|]+\|", b.strip()):
                    continue
                cells = [c.strip() for c in b.strip().strip("|").split("|")]
                rows.append(cells)
            if rows:
                cols = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=cols); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
                widths = []
                for c in range(cols):
                    m = 0
                    for r in rows:
                        txt = r[c] if c < len(r) else ""
                        # 中文按 1 字宽、ASCII 按 0.55 字宽估算显示宽度
                        w = sum(1.0 if ord(ch) > 0x2E80 else 0.55 for ch in txt)
                        m = max(m, w)
                    widths.append(max(1.6, min(5.2, m * 0.26)))
                total = sum(widths)
                avail = 16.6
                widths = [w / total * avail for w in widths]
                set_table_fixed(t, widths)
                for ri, r in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(cols):
                        txt = r[ci] if ci < len(r) else ""
                        cell = cells[ci]; cell.width = Cm(widths[ci])
                        par = cell.paragraphs[0]; par.paragraph_format.space_after = Pt(1)
                        add_inline(par, txt, size=8.5, base_bold=(ri == 0))
                        if ri == 0:
                            shade(cell._tc, "DEEAF6")
                doc.add_paragraph()
            continue
        if not line.strip() or line.strip() == "---":
            i += 1; continue
        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(10)
                r = p.add_run(text); set_fonts(r, "Times New Roman", "微软雅黑", 18, True, "1F3864")
                p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r2 = p2.add_run("架构设计规范文档  |  CampusVenue  |  V1.0"); set_fonts(r2, "Times New Roman", "微软雅黑", 10.5, False, "595959")
            else:
                sizes = {2: 15, 3: 12.5, 4: 11.5}
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12 if level == 2 else 8); p.paragraph_format.space_after = Pt(4)
                r = p.add_run(text); set_fonts(r, "Times New Roman", "微软雅黑", sizes[level], True, "1F3864" if level == 2 else "2E5496")
                if level == 2:
                    shade(p._p, "EDF2F9")
            if not img_inserted and text.startswith("4.3"):
                doc.add_paragraph()
                pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(diagram, width=Cm(16.4))
                cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rc = cap.add_run("图 4-1  总体架构与部署拓扑"); set_fonts(rc, "Times New Roman", "微软雅黑", 9, False, "595959")
                img_inserted = True
            i += 1; continue
        # 引用
        if line.startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, line.lstrip("> ").strip(), size=10.5, base_bold=True, color="1F3864")
            i += 1; continue
        # 复选清单
        m = re.match(r"^-\s*\[( |x)\]\s*(.*)$", line)
        if m:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run(("☑ " if m.group(1) == "x" else "☐ ") + m.group(2))
            set_fonts(r, "Times New Roman", "宋体", 10.5)
            i += 1; continue
        # 列表
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Number" if m.group(2)[0].isdigit() else "List Bullet")
            p.paragraph_format.left_indent = Cm(0.75 + indent * 0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(3), size=10.5)
            i += 1; continue
        # 正文
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(0)
        add_inline(p, line.strip(), size=10.5)
        i += 1

    doc.save(DOCX)
    return DOCX

if __name__ == "__main__":
    img = build_diagram()
    out = build_docx(img)
    print("diagram:", img, os.path.getsize(img))
    print("docx:", out, os.path.getsize(out))
